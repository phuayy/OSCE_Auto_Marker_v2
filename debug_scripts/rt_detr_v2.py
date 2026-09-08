#!/usr/bin/env python3
"""N-process time-chunk parallel variant of ``rt_detr.py`` (benchmark harness).

Splits one video into N contiguous time-chunks (aligned to the sample grid),
spawns N worker PROCESSES that each load their own RT-DETR model and detect
people over their chunk, then merges the per-chunk person-count arrays back
into one global timeline and runs the SAME tolerant N-of-M segmenter as
``rt_detr.py`` (imported, not copied) to decide session clips (>= 2 people)
vs idle (1/0 people). A session that straddles a chunk boundary is still
detected because segmentation runs ONCE over the merged counts, never
per-chunk.

Purpose: a test run to compare single-process (rt_detr.py) vs multi-process
wall-clock on the same video/params. Expectations on a single 4 GB GPU
(RTX 3050): ~no GPU speedup (CUDA serializes across processes), 2x model
VRAM, real OOM risk — the per-process OOM batch-halving/CPU-fallback from the
production detector is reused so runs degrade instead of dying. On CPU
(``--device cpu``) the N processes ARE genuinely parallel.

Chunk alignment: chunks are split on SAMPLE INDEX (sample k is at k/sample_fps
seconds). Each worker seeks with ``ffmpeg -ss <k0*dt>`` and samples with the
same ``fps=`` filter, so its frames land on the global grid at k0, k0+1, ...
Workers cap their output at the planned sample count; a chunk that decodes
short (EOF rounding) is padded by repeating its last count so the global
timeline stays index-aligned.

First run note: warm the HuggingFace model cache with a single-worker or
rt_detr.py run first — N cold workers racing the same download is untested.

Run (same interpreter as rt_detr.py):

    uv run python debug_scripts/rt_detr_v2.py \
        --video "C:/Users/.../Common Cold_Session 1.mp4" --workers 2

    python debug_scripts/rt_detr_v2.py --self-check   # split-math sanity check
"""

from __future__ import annotations

import argparse
import json
import multiprocessing as mp
import subprocess
import sys
import time
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from typing import Any

# Reuse rt_detr.py (which itself puts scripts/ on sys.path and imports the
# production detect_human_segments module). Single source of truth for the
# detector, box drawing, tolerant segmenter, and the xlsx report.
_HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(_HERE))
import rt_detr as rt  # noqa: E402

dhs = rt.dhs

DEFAULT_OUT_ROOT = rt.DEFAULT_OUT_ROOT


# Chunk planning + chunked frame decoding were promoted into production
# (scripts/detect_human_segments.py, 2026-07-17) — aliased here so this
# harness stays a thin benchmark wrapper around the same code paths.
split_sample_indices = dhs.split_sample_indices
iter_chunk_frames = dhs.iter_chunk_frames
self_check = dhs.self_check


# ---------------------------------------------------------------------------
# Worker (runs in a spawned child process — must be top-level & picklable I/O)
# ---------------------------------------------------------------------------


def run_chunk(task: dict) -> dict:
    """Decode + detect one chunk. Returns per-sample counts (small) — annotated
    PNGs are written directly from the worker (global frame index in the
    filename), so no frames cross the process boundary."""
    wall_started = time.perf_counter()
    wid = int(task["worker"])
    start_index = int(task["start_index"])
    end_index = int(task["end_index"])
    expected = end_index - start_index
    sample_dt = float(task["sample_dt"])

    def wlog(message: str) -> None:
        dhs.log(f"[worker {wid}] {message}")

    det_cfg = dhs.DetectorConfig(
        model_id=task["model_id"],
        device=task["device"],
        confidence=task["confidence"],
        batch_size=task["batch_size"],
        decode_width=task["decode_width"],
    )
    annotations_dir = Path(task["annotations_dir"]) if task["annotations_dir"] else None

    load_started = time.perf_counter()
    counter = dhs.PersonCounter(det_cfg)
    counter.load()
    load_seconds = time.perf_counter() - load_started
    wlog(
        f"model ready in {load_seconds:.1f}s ({counter.device}); "
        f"samples {start_index}..{end_index - 1} (t={start_index * sample_dt:.0f}s..)"
    )

    counts: list[int] = []
    batch: list[Any] = []
    decode_started = time.perf_counter()
    inference_seconds = 0.0
    annotate_seconds = 0.0

    def flush(batch_frames: list[Any]) -> None:
        nonlocal inference_seconds, annotate_seconds
        t0 = time.perf_counter()
        boxes = rt.detect_person_boxes(counter, batch_frames)
        inference_seconds += time.perf_counter() - t0
        local_start = len(counts)
        for frame_boxes in boxes:
            counts.append(len(frame_boxes))
        if annotations_dir is not None:
            t1 = time.perf_counter()
            for j, (frm, frame_boxes) in enumerate(zip(batch_frames, boxes)):
                global_idx = start_index + local_start + j
                t_s = global_idx * sample_dt
                out = annotations_dir / f"frame_{global_idx:06d}_t{int(t_s):06d}s.png"
                rt.annotate_and_save(frm, frame_boxes, out)
            annotate_seconds += time.perf_counter() - t1

    try:
        for frame in iter_chunk_frames(
            task["ffmpeg_bin"],
            Path(task["video_path"]),
            start_seconds=start_index * sample_dt,
            duration_seconds=expected * sample_dt,
            sample_fps=task["sample_fps"],
            out_width=task["out_width"],
            out_height=task["out_height"],
        ):
            batch.append(frame)
            if len(batch) >= det_cfg.batch_size:
                flush(batch)
                batch = []
                if len(counts) % 120 < det_cfg.batch_size:
                    wlog(f"analysed {len(counts)}/{expected} samples")
            if len(counts) + len(batch) >= expected:
                break  # planned sample count reached — stop decoding
        if batch:
            flush(batch)
    finally:
        counter.close()

    decode_seconds = time.perf_counter() - decode_started - inference_seconds - annotate_seconds
    counts = counts[:expected]  # safety cap; padding for short chunks is the parent's job
    wlog(f"done: {len(counts)}/{expected} samples in {time.perf_counter() - wall_started:.1f}s")
    return {
        "chunk_index": int(task["chunk_index"]),
        "start_index": start_index,
        "end_index": end_index,
        "counts": counts,
        "load_seconds": round(load_seconds, 3),
        "decode_seconds": round(decode_seconds, 3),
        "inference_seconds": round(inference_seconds, 3),
        "annotate_seconds": round(annotate_seconds, 3),
        "wall_seconds": round(time.perf_counter() - wall_started, 3),
        "device": counter.device,
        "model": counter.resolved_model_id,
    }


# ---------------------------------------------------------------------------
# Main (parent): plan chunks -> spawn workers -> merge -> segment -> report
# ---------------------------------------------------------------------------


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="RT-DETR human detection, N parallel processes over time-chunks (benchmark vs rt_detr.py)."
    )
    p.add_argument("--video", default=rt.DEFAULT_VIDEO)
    p.add_argument("--workers", type=int, default=2, help="Parallel detector processes (N); each loads its own model")
    p.add_argument("--annotations-dir", default=str(Path(DEFAULT_OUT_ROOT) / "annotations_v2"))
    p.add_argument("--report", default=str(Path(DEFAULT_OUT_ROOT) / "rt_detr_v2_report.xlsx"))
    p.add_argument("--sample-fps", type=float, default=1.0)
    p.add_argument("--min-people", type=int, default=2)
    p.add_argument("--start-after-seconds", type=float, default=50.0)
    p.add_argument("--end-after-seconds", type=float, default=50.0)
    p.add_argument("--flicker-tolerance-seconds", type=float, default=10.0)
    p.add_argument("--median-window", type=int, default=3, help="Unused by the tolerant segmenter; kept for CLI parity")
    p.add_argument("--start-offset", type=float, default=0.0)
    p.add_argument("--end-offset", type=float, default=2.0)
    p.add_argument("--min-clip-seconds", type=float, default=0.5)
    p.add_argument("--confidence", type=float, default=0.7)
    p.add_argument("--model", default=dhs.DEFAULT_MODEL)
    p.add_argument("--device", choices=["auto", "cuda", "cpu"], default="auto")
    p.add_argument("--batch-size", type=int, default=8, help="Per-worker batch (lower to 4 on 4GB VRAM with 2 workers)")
    p.add_argument("--decode-width", type=int, default=640)
    p.add_argument("--max-frames", type=int, default=0, help="Cap TOTAL sampled frames (0 = all) — for quick tests")
    p.add_argument("--no-annotate", action="store_true", help="Skip writing PNGs (counts + report only)")
    p.add_argument("--self-check", action="store_true", help="Run split-math asserts and exit")
    return p.parse_args()


def main() -> int:
    script_started = time.perf_counter()
    args = parse_args()
    if args.self_check:
        self_check()
        return 0

    video_path = Path(args.video).expanduser().resolve()
    if not video_path.exists():
        raise FileNotFoundError(f"Video not found: {video_path}")

    annotations_dir = Path(args.annotations_dir).expanduser()
    annotations_dir.mkdir(parents=True, exist_ok=True)
    report_path = Path(args.report).expanduser()

    seg_cfg = dhs.SegmenterConfig(
        sample_fps=max(0.1, args.sample_fps),
        min_people=max(1, args.min_people),
        start_after_seconds=max(0.0, args.start_after_seconds),
        end_after_seconds=max(0.0, args.end_after_seconds),
        flicker_tolerance_seconds=max(0.0, args.flicker_tolerance_seconds),
        start_offset_seconds=max(0.0, args.start_offset),
        end_offset_seconds=max(0.0, args.end_offset),
        min_clip_seconds=max(0.0, args.min_clip_seconds),
    )

    ffmpeg_bin = dhs.resolve_binary("FFMPEG_BIN", "ffmpeg")
    ffprobe_bin = dhs.resolve_binary("FFPROBE_BIN", "ffprobe")
    src_w, src_h = dhs.probe_video_dimensions(ffprobe_bin, video_path)
    out_w = min(max(160, args.decode_width), src_w)
    out_h = max(2, round(src_h * out_w / src_w / 2) * 2)
    out_w = max(2, out_w // 2 * 2)

    probe = subprocess.run(
        [ffprobe_bin, "-v", "error", "-show_entries", "format=duration", "-of", "json", str(video_path)],
        capture_output=True, text=True, check=True,
    )
    video_duration = float(json.loads(probe.stdout)["format"]["duration"])

    total_samples = int(video_duration * seg_cfg.sample_fps) + 1
    if args.max_frames:
        total_samples = min(total_samples, args.max_frames)
    spans = split_sample_indices(total_samples, max(1, args.workers))
    if not spans:
        raise RuntimeError("Video too short to sample any frames.")

    dhs.log(
        f"v2: {len(spans)} worker(s) over {total_samples} samples of {video_path.name} "
        f"@ {seg_cfg.sample_fps} fps ({src_w}x{src_h} -> {out_w}x{out_h}); "
        f"duration {rt._fmt_ts(video_duration)}."
    )
    for i, (a, b) in enumerate(spans):
        dhs.log(f"  chunk {i}: samples {a}..{b - 1}  (t={a * seg_cfg.sample_dt:.0f}s..{(b - 1) * seg_cfg.sample_dt:.0f}s)")

    tasks = [
        {
            "chunk_index": i,
            "worker": i,
            "start_index": a,
            "end_index": b,
            "sample_dt": seg_cfg.sample_dt,
            "sample_fps": seg_cfg.sample_fps,
            "video_path": str(video_path),
            "ffmpeg_bin": ffmpeg_bin,
            "out_width": out_w,
            "out_height": out_h,
            "model_id": args.model,
            "device": args.device,
            "confidence": min(0.99, max(0.05, args.confidence)),
            "batch_size": max(1, args.batch_size),
            "decode_width": max(160, args.decode_width),
            "annotations_dir": "" if args.no_annotate else str(annotations_dir),
        }
        for i, (a, b) in enumerate(spans)
    ]

    pool_started = time.perf_counter()
    ctx = mp.get_context("spawn")
    with ProcessPoolExecutor(max_workers=len(tasks), mp_context=ctx) as pool:
        results = list(pool.map(run_chunk, tasks))  # map preserves chunk order
    # Wall-clock of the parallel detection phase only (spawn+load+decode+infer+
    # annotate across all N processes running concurrently) — NOT the script total.
    pool_wall_seconds = time.perf_counter() - pool_started

    # --- Merge: concatenate per-chunk counts onto the global sample grid ---
    counts: list[int] = []
    padded_indices: set[int] = set()
    for res in results:
        expected = res["end_index"] - res["start_index"]
        chunk_counts = list(res["counts"])
        if len(chunk_counts) < expected:
            short_by = expected - len(chunk_counts)
            pad_value = chunk_counts[-1] if chunk_counts else 0
            dhs.log(f"chunk {res['chunk_index']} decoded {short_by} sample(s) short — padding with {pad_value}.")
            for k in range(short_by):
                padded_indices.add(res["start_index"] + len(chunk_counts) + k)
            chunk_counts.extend([pad_value] * short_by)
        res["padded_samples"] = expected - len(res["counts"])
        counts.extend(chunk_counts)

    if not counts:
        raise RuntimeError("No frames decoded — is the video readable?")

    histogram: dict[str, int] = {}
    for c in counts:
        histogram[str(c)] = histogram.get(str(c), 0) + 1
    dhs.log(f"Person-count histogram: {histogram}")

    # --- Rules: ONE global pass of the tolerant segmenter (2 people = session,
    # <2 = idle), identical to rt_detr.py — sessions may straddle chunk seams.
    segment_started = time.perf_counter()
    clip_ranges, transitions = rt.build_tolerant_session_ranges(
        counts,
        video_duration,
        min_people=seg_cfg.min_people,
        sample_dt=seg_cfg.sample_dt,
        start_window_samples=seg_cfg.seconds_to_samples(seg_cfg.start_after_seconds),
        start_tolerance_samples=seg_cfg.seconds_to_samples(seg_cfg.flicker_tolerance_seconds),
        end_window_samples=seg_cfg.seconds_to_samples(seg_cfg.end_after_seconds),
        end_tolerance_samples=seg_cfg.seconds_to_samples(seg_cfg.flicker_tolerance_seconds),
        start_offset_seconds=seg_cfg.start_offset_seconds,
        end_offset_seconds=seg_cfg.end_offset_seconds,
        min_clip_seconds=seg_cfg.min_clip_seconds,
    )
    segment_seconds = time.perf_counter() - segment_started
    dhs.log(f"Detected {len(clip_ranges)} student clip(s).")

    # --- Frame rows (same shape as rt_detr.py) ---
    dt = seg_cfg.sample_dt
    frame_rows: list[dict] = []
    for idx, count in enumerate(counts):
        t_s = idx * dt
        clip_idx = ""
        in_session = False
        for clip in clip_ranges:
            if clip["start"] <= t_s <= clip["end"]:
                in_session = True
                clip_idx = clip.get("student_index", "")
                break
        if args.no_annotate:
            fname = ""
        elif idx in padded_indices:
            fname = "(padded — no frame decoded)"
        else:
            fname = f"frame_{idx:06d}_t{int(t_s):06d}s.png"
        frame_rows.append({
            "frame_index": idx, "time_s": t_s, "person_count": count,
            "in_session": in_session, "clip_index": clip_idx, "file": fname,
        })

    # --- Aggregates across the N worker processes ---
    # "sum" = total compute cost (apples-to-apples vs rt_detr.py's single-process
    # time); "max" = the slowest worker, i.e. what actually gates wall-clock
    # since the workers run concurrently.
    total_load_seconds = sum(r["load_seconds"] for r in results)
    max_load_seconds = max(r["load_seconds"] for r in results)
    total_decode_seconds = sum(r["decode_seconds"] for r in results)
    total_inference_seconds = sum(r["inference_seconds"] for r in results)
    max_inference_seconds = max(r["inference_seconds"] for r in results)
    total_annotate_seconds = sum(r["annotate_seconds"] for r in results)

    # "Start to end" total — the real number to compare against rt_detr.py's
    # TOTAL (probe + spawn + N-worker detection + merge + segmentation).
    total_seconds = time.perf_counter() - script_started

    rt.write_report(
        report_path,
        video_path=video_path,
        video_duration=video_duration,
        seg_cfg=seg_cfg,
        resolved_model=results[0]["model"],
        device=", ".join(sorted({r["device"] for r in results})),
        confidence=min(0.99, max(0.05, args.confidence)),
        resolution=f"{out_w}x{out_h}",
        counts=counts,
        frame_rows=frame_rows,
        clip_ranges=clip_ranges,
        transitions=transitions,
        histogram=histogram,
        load_seconds=total_load_seconds,
        decode_seconds=total_decode_seconds,
        inference_seconds=total_inference_seconds,
        annotate_seconds=total_annotate_seconds,
        segment_seconds=segment_seconds,
        total_seconds=total_seconds,
        annotations_dir=annotations_dir,
    )

    # --- Extra "Workers" sheet: the per-process benchmark breakdown ---
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(report_path)
    ws = wb.create_sheet("Workers")
    ws.append([
        "worker", "start_sample", "end_sample", "samples", "padded",
        "load_s", "decode_s", "inference_s", "annotate_s", "wall_s", "device", "model",
    ])
    for res in results:
        ws.append([
            res["chunk_index"], res["start_index"], res["end_index"] - 1,
            res["end_index"] - res["start_index"], res["padded_samples"],
            res["load_seconds"], res["decode_seconds"], res["inference_seconds"],
            res["annotate_seconds"], res["wall_seconds"], res["device"], res["model"],
        ])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Workers (N)", len(results)])
    ws.append(["Sum of load time across workers (s) — total compute cost", round(total_load_seconds, 1)])
    ws.append(["Max load time across workers (s) — parallel wall proxy", round(max_load_seconds, 1)])
    ws.append(["Sum of inference time across workers (s) — total inference", round(total_inference_seconds, 1)])
    ws.append(["Max inference time across workers (s) — parallel wall proxy", round(max_inference_seconds, 1)])
    ws.append(["Detection-phase wall-clock (s) — pool.map, parallel", round(pool_wall_seconds, 1)])
    ws.append(["Segmentation/rules time (s)", round(segment_seconds, 1)])
    ws.append(["TOTAL SCRIPT wall-clock (s) — compare vs rt_detr.py", round(total_seconds, 1)])
    wb.save(report_path)

    # --- Console summary ---
    print("\n" + "=" * 60)
    print(f"  RT-DETR v2 ({len(results)} processes) — {video_path.name}")
    print("=" * 60)
    print(f"  Sampled frames : {len(counts)} ({len(padded_indices)} padded)")
    print(f"  Annotated PNGs : {'(skipped)' if args.no_annotate else len(counts) - len(padded_indices)} -> {annotations_dir}")
    print(f"  Clips detected : {len(clip_ranges)}")
    for clip in clip_ranges:
        sf, ef = round(clip["start"] / dt), round(clip["end"] / dt)
        print(
            f"    #{clip['student_index']:>2}  {rt._fmt_ts(clip['start'])}–{rt._fmt_ts(clip['end'])}  "
            f"(frames {sf}–{ef}, {clip['end'] - clip['start']:.0f}s)"
        )
    print(f"  Model load     : sum {total_load_seconds:.1f}s across {len(results)}  (max/parallel-wall {max_load_seconds:.1f}s)")
    print(f"  Decode         : sum {total_decode_seconds:.1f}s across {len(results)}")
    print(f"  Inference      : sum {total_inference_seconds:.1f}s across {len(results)}  (max/parallel-wall {max_inference_seconds:.1f}s)")
    print(f"  Detection pool : {pool_wall_seconds:.1f}s wall-clock (parallel; spawn+load+decode+infer+annotate)")
    print(f"  Segmentation   : {segment_seconds:.1f}s  (clip detection rules)")
    print(f"  TOTAL          : {total_seconds:.1f}s  <-- compare against rt_detr.py TOTAL")
    for res in results:
        print(
            f"    worker {res['chunk_index']}: load {res['load_seconds']:.1f}s  "
            f"decode {res['decode_seconds']:.1f}s  infer {res['inference_seconds']:.1f}s  "
            f"annotate {res['annotate_seconds']:.1f}s  wall {res['wall_seconds']:.1f}s  [{res['device']}]"
        )
    print(f"  Report         : {report_path}")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
